import sys

from pathlib import Path
from html import escape
from openpyxl import load_workbook

if len(sys.argv) != 2:
    raise SystemExit("변환할 파일 경로를 인자로 전달해주세요.")

input_path = Path(sys.argv[1]).expanduser().resolve()

# 문서 이름과 확장자로 결과 폴더 구분
extension = input_path.suffix.lower().lstrip(".")

output_dir = (
    Path(__file__).resolve().parent
    / "html_output"
    / f"{input_path.stem}_{extension}"
)

output_dir.mkdir(parents=True, exist_ok=True)

output_path = output_dir / "index.html"

if not input_path.is_file():
    raise FileNotFoundError(f"파일을 찾을 수 없습니다: {input_path}")

output_dir.mkdir(parents=True, exist_ok=True)

# 수식 자체 대신 파일에 저장된 계산 결과를 읽기
workbook = load_workbook(
    input_path,
    data_only=True,
    read_only=True,
)

try:
    sheets_html = []

    # 모든 시트를 순서대로 처리
    for sheet in workbook.worksheets:
        rows_html = []

        for row in sheet.iter_rows(values_only=True):
            cells_html = []

            for value in row:
                text = "" if value is None else str(value)
                cells_html.append(f"<td>{escape(text)}</td>")

            rows_html.append("<tr>" + "".join(cells_html) + "</tr>")

        table_html = "\n".join(rows_html)
        sheet_name = escape(sheet.title)

        # 시트 이름과 표를 하나의 구역으로 묶기
        sheets_html.append(
            f"<section>"
            f"<h2>{sheet_name}</h2>"
            f"<table>{table_html}</table>"
            f"</section>"
        )

    # 모든 시트의 HTML 합치기
    all_sheets_html = "\n".join(sheets_html)

finally:
    workbook.close()
    
html_content = f"""<!DOCTYPE html>
<html lang="ko">
<head>
    <meta charset="utf-8">
    <title>XLSX 변환 결과</title>
    <style>
        body {{
            padding: 24px;
            font-family: "맑은 고딕", sans-serif;
        }}
        section {{
            margin-bottom: 40px;
        }}
        table {{
            border-collapse: collapse;
        }}
        td {{
            border: 1px solid #aaa;
            padding: 8px 12px;
            white-space: pre-wrap;
        }}
    </style>
</head>
<body>
    <h1>Excel 변환 결과</h1>
    {all_sheets_html}
</body>
</html>
"""

# HTML 문자열을 실제 파일로 저장
output_path.write_text(html_content, encoding="utf-8")

# 터미널에 완료 메시지 출력
print(f"변환 완료: {output_path}")